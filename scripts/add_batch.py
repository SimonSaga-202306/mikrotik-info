#!/usr/bin/env python3
"""Append a batch of consultant rows to the reverse-signal list.

Usage: python3 scripts/add_batch.py batch.json

Each item: {"name","country","city","site","cert","sku","prio","note","contact"}
plus optional "angle", "type", "signal".
"""
import csv, json, re, sys
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

CSV = "客户名单_反向信号_2026-09.csv"
XLSX = "MikroTik_vs_SmartByte_替代对照表.xlsx"
HDR = list(csv.reader(open(CSV, encoding="utf-8-sig")))[0]

DEF_ANGLE = ("MikroTik 认证顾问：自己做方案/集成/运维并自行采购设备 → "
             "用 ERPS/IGMP/VLAN/CLI/MIB 这类技术语言切入，不打价格战")

SHEET_ORDER = ["参数", "替代对照表", "竞品全景", "Marvell品牌全景", "A类客户速查",
               "反向信号名单", "欧洲街价锚点", "重点客户配货表", "MikroTik全产品线",
               "合规文件清单", "开发系统-节奏与KPI", "MikroTik出口地(Comtrade2024)",
               "落地价与关税", "MikroTik关键事实(实测)", "SmartByte关键SKU", "原Excel数据修正"]


def norm(s):
    return re.sub(r'[^a-z0-9]+', '', (s or "").lower())


def main():
    items = json.load(open(sys.argv[1], encoding="utf-8"))
    rows = list(csv.reader(open(CSV, encoding="utf-8-sig")))
    old = rows[1:]
    existing = set()
    for r in old:
        existing.add(norm(r[0]))
        if r[3]:
            existing.add(norm(r[3].split('/')[0]))

    added, skipped = [], []
    for it in items:
        r = [it["name"], it["country"], it.get("city", ""), it.get("site", ""),
             it.get("type", "集成商/顾问（MikroTik 认证）"),
             it.get("signal", "MikroTik 认证顾问 " + it.get("cert", "")),
             "否", it.get("angle", DEF_ANGLE), it.get("sku", ""),
             it.get("contact", ""), it.get("prio", "★★"),
             "未触达", "2026-09-22", "", "", "", it.get("note", "")]
        assert len(r) == 17, it["name"]
        k = norm(r[0])
        k2 = norm(r[3].split('/')[0]) if r[3] else ""
        if k in existing or (k2 and k2 in existing):
            skipped.append(r[0])
            continue
        existing.add(k)
        if k2:
            existing.add(k2)
        added.append(r)

    allr = old + added
    with open(CSV, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(HDR)
        w.writerows(allr)

    wb = load_workbook(XLSX)
    hf = PatternFill("solid", fgColor="1F4E78")
    if "反向信号名单" in wb.sheetnames:
        del wb["反向信号名单"]
    ws = wb.create_sheet("反向信号名单")
    ws.append(HDR)
    for r in allr:
        ws.append(r)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = hf
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for i, w_ in enumerate([30, 14, 14, 24, 18, 32, 12, 34, 24, 30, 8, 10, 18,
                            12, 12, 12, 30], 1):
        ws.column_dimensions[get_column_letter(i)].width = w_
    ws.freeze_panes = "B2"
    wb._sheets = [wb[n] for n in SHEET_ORDER if n in wb.sheetnames]
    wb.save(XLSX)

    print("added:", len(added), "| skipped:", skipped)
    print("total:", len(allr))
    for r in skipped:
        print("  CHECK-SKIP:", r)


if __name__ == "__main__":
    main()
